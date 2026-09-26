
from pathlib import Path
import random
from copy import copy

from openpyxl import Workbook


# ============================================================
# 1. PROJECT SETUP
# ============================================================

random.seed(42)

PROJECT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_DIR / "data"

DATA_DIR.mkdir(exist_ok=True)


# ============================================================
# 2. COMPANY STRUCTURE
# ============================================================

ENTITIES = [
    "Austria",
    "Germany",
    "France",
    "Switzerland"
]


ACCOUNTS = [
    ("100000", "Cash", "Balance Sheet"),
    ("110000", "Accounts Receivable", "Balance Sheet"),
    ("120000", "Inventory", "Balance Sheet"),
    ("130000", "Prepayments", "Balance Sheet"),
    ("200000", "Accounts Payable", "Balance Sheet"),
    ("210000", "Accrued Liabilities", "Balance Sheet"),
    ("300000", "Share Capital", "Balance Sheet"),
    ("310000", "Retained Earnings", "Balance Sheet"),
    ("400000", "Revenue", "P&L"),
    ("410000", "Other Operating Income", "P&L"),
    ("500000", "Materials", "P&L"),
    ("510000", "Production Costs", "P&L"),
    ("600000", "Personnel Costs", "P&L"),
    ("610000", "Utilities", "P&L"),
    ("620000", "Consulting", "P&L"),
    ("630000", "IT Services", "P&L"),
    ("640000", "Insurance", "P&L"),
    ("650000", "Software Licences", "P&L"),
    ("660000", "Travel & Expenses", "P&L"),
    ("670000", "Marketing", "P&L"),
]


# ============================================================
# 3. CREATE TRIAL BALANCE
# ============================================================

tb_rows = []


for entity in ENTITIES:

    debit_accounts = [
        "100000",
        "110000",
        "120000",
        "130000",
        "500000",
        "510000",
        "600000",
        "610000",
        "620000",
        "630000",
        "640000",
        "650000",
        "660000",
        "670000"
    ]

    credit_accounts = [
        "200000",
        "210000",
        "300000",
        "400000",
        "410000"
    ]

    balances = {}


    # Generate debit balances

    for account in debit_accounts:

        balances[account] = round(
            random.uniform(5_000, 180_000),
            2
        )


    # Generate credit balances

    for account in credit_accounts:

        balances[account] = round(
            random.uniform(10_000, 300_000),
            2
        )


    # Calculate retained earnings
    # as the balancing figure

    total_debit = sum(
        balances[account]
        for account in debit_accounts
    )

    total_credit = sum(
        balances[account]
        for account in credit_accounts
    )

    balances["310000"] = round(
        total_debit - total_credit,
        2
    )


    # Add accounts to the dataset

    for account, account_name, category in ACCOUNTS:

        amount = balances.get(account, 0)


        if account in debit_accounts:

            debit = amount
            credit = 0


        elif account in credit_accounts or account == "310000":

            debit = 0
            credit = amount


        else:

            debit = 0
            credit = 0


        tb_rows.append([
            account,
            account_name,
            entity,
            debit,
            credit,
            category
        ])


# ============================================================
# 4. SAVE TRIAL BALANCE
# ============================================================

tb_file_path = DATA_DIR / "Trial_Balance_December_2025.xlsx"

tb_workbook = Workbook()

tb_worksheet = tb_workbook.active
tb_worksheet.title = "Trial Balance"


tb_headers = [
    "Account",
    "Account Name",
    "Entity",
    "Debit",
    "Credit",
    "Category"
]

tb_worksheet.append(tb_headers)


for row in tb_rows:

    tb_worksheet.append(row)


# Formatting

for cell in tb_worksheet[1]:

    new_font = copy(cell.font)
    new_font.bold = True
    cell.font = new_font


tb_worksheet.freeze_panes = "A2"
tb_worksheet.auto_filter.ref = tb_worksheet.dimensions


tb_column_widths = {
    "A": 14,
    "B": 25,
    "C": 15,
    "D": 16,
    "E": 16,
    "F": 18
}


for column, width in tb_column_widths.items():

    tb_worksheet.column_dimensions[column].width = width


tb_workbook.save(tb_file_path)


print("Trial Balance successfully created.")
print(f"File: {tb_file_path}")
print(f"Rows created: {len(tb_rows)}")


# ============================================================
# 5. CREATE ACCRUALS
# ============================================================

ACCRUAL_ACCOUNTS = [
    "600000",
    "610000",
    "620000",
    "630000",
    "660000"
]


ACCRUAL_DESCRIPTIONS = [
    "Personnel bonus accrual",
    "Utilities December",
    "Consulting services",
    "IT services",
    "Travel expenses",
    "Production services"
]


accrual_rows = []


for i in range(1, 36):

    entity = random.choice(ENTITIES)

    account = random.choice(
        ACCRUAL_ACCOUNTS
    )

    description = random.choice(
        ACCRUAL_DESCRIPTIONS
    )

    amount = round(
        random.uniform(2_500, 35_000),
        2
    )

    accrual_rows.append([
        f"ACC{i:03d}",
        account,
        entity,
        description,
        amount,
        "2025-12-31",
        "Open"
    ])


# ============================================================
# 6. INTENTIONAL ACCRUAL ANOMALY
# ============================================================

# Account 699999 does not exist
# in the Trial Balance.

accrual_rows.append([
    "ACC999",
    "699999",
    "Austria",
    "Unmapped closing accrual",
    8_500.00,
    "2025-12-31",
    "Open"
])


# ============================================================
# 7. SAVE ACCRUALS
# ============================================================

accrual_file_path = DATA_DIR / "Accruals_December_2025.xlsx"

accrual_workbook = Workbook()

accrual_worksheet = accrual_workbook.active
accrual_worksheet.title = "Accruals"


accrual_headers = [
    "Accrual ID",
    "Account",
    "Entity",
    "Description",
    "Amount",
    "Posting Date",
    "Status"
]

accrual_worksheet.append(
    accrual_headers
)


for row in accrual_rows:

    accrual_worksheet.append(row)


# Formatting

for cell in accrual_worksheet[1]:

    new_font = copy(cell.font)
    new_font.bold = True
    cell.font = new_font


accrual_worksheet.freeze_panes = "A2"
accrual_worksheet.auto_filter.ref = accrual_worksheet.dimensions


accrual_column_widths = {
    "A": 14,
    "B": 14,
    "C": 15,
    "D": 30,
    "E": 16,
    "F": 16,
    "G": 12
}


for column, width in accrual_column_widths.items():

    accrual_worksheet.column_dimensions[column].width = width


accrual_workbook.save(
    accrual_file_path
)


print("Accruals successfully created.")
print(f"File: {accrual_file_path}")
print(f"Rows created: {len(accrual_rows)}")

# ============================================================
# 8. CREATE PREPAYMENTS
# ============================================================

PREPAYMENT_ACCOUNTS = [
    "130000"
]

PREPAYMENT_DESCRIPTIONS = [
    "Annual insurance",
    "Software licence",
    "Annual maintenance contract",
    "Rent",
    "Service contract",
    "Subscription"
]

prepayment_rows = []

for i in range(1, 21):

    entity = random.choice(ENTITIES)

    description = random.choice(
        PREPAYMENT_DESCRIPTIONS
    )

    amount = round(
        random.uniform(1_000, 20_000),
        2
    )

    expense_account = random.choice([
        "640000",
        "650000",
        "620000",
        "630000"
    ])

    prepayment_rows.append([
        f"PREP{i:03d}",
        "130000",
        expense_account,
        entity,
        description,
        amount,
        "2025-12-31",
        "Open"
    ])


# ============================================================
# 9. INTENTIONAL PREPAYMENT ANOMALY
# ============================================================

# Account 799998 does not exist
# in the Trial Balance.

prepayment_rows.append([
    "PREP999",
    "130000",
    "799998",
    "Germany",
    "Unmapped prepaid expense",
    6_750.00,
    "2025-12-31",
    "Open"
])


# ============================================================
# 10. SAVE PREPAYMENTS
# ============================================================

prepayment_file_path = (
    DATA_DIR / "Prepayments_December_2025.xlsx"
)

prepayment_workbook = Workbook()

prepayment_worksheet = prepayment_workbook.active
prepayment_worksheet.title = "Prepayments"


prepayment_headers = [
    "Prepayment ID",
    "Balance Sheet Account",
    "Expense Account",
    "Entity",
    "Description",
    "Amount",
    "Posting Date",
    "Status"
]

prepayment_worksheet.append(
    prepayment_headers
)

for row in prepayment_rows:
    prepayment_worksheet.append(row)


# Formatting

for cell in prepayment_worksheet[1]:

    new_font = copy(cell.font)
    new_font.bold = True
    cell.font = new_font

prepayment_worksheet.freeze_panes = "A2"
prepayment_worksheet.auto_filter.ref = (
    prepayment_worksheet.dimensions
)


prepayment_column_widths = {
    "A": 16,
    "B": 22,
    "C": 18,
    "D": 15,
    "E": 30,
    "F": 16,
    "G": 16,
    "H": 12
}

for column, width in prepayment_column_widths.items():
    prepayment_worksheet.column_dimensions[column].width = width


prepayment_workbook.save(
    prepayment_file_path
)

print("Prepayments successfully created.")
print(f"File: {prepayment_file_path}")
print(f"Rows created: {len(prepayment_rows)}")

# ============================================================
# 11. CREATE INVOICES
# ============================================================

INVOICE_ACCOUNTS = [
    "500000",
    "510000",
    "620000",
    "630000",
    "640000",
    "650000",
    "660000",
    "670000"
]

VENDORS = [
    "Office Supplies GmbH",
    "Industrial Services AG",
    "IT Solutions GmbH",
    "Facility Management SA",
    "Travel Services GmbH",
    "Energy Services AG",
    "Consulting Partners GmbH",
    "Software Solutions AG"
]

invoice_rows = []

for i in range(1, 31):

    entity = random.choice(ENTITIES)

    vendor = random.choice(VENDORS)

    account = random.choice(
        INVOICE_ACCOUNTS
    )

    amount = round(
        random.uniform(1_000, 50_000),
        2
    )

    invoice_rows.append([
        f"INV{i:04d}",
        vendor,
        account,
        entity,
        f"Invoice {i:04d}",
        amount,
        "2025-12-31",
        "Posted"
    ])


# ============================================================
# 12. INTENTIONAL DUPLICATE INVOICE
# ============================================================

# Same vendor, account, entity, description and amount.
# Only the invoice ID is different.

duplicate_invoice = invoice_rows[4].copy()

duplicate_invoice[0] = "INV9999"

invoice_rows.append(
    duplicate_invoice
)


# ============================================================
# 13. INTENTIONAL UNMAPPED ACCOUNT
# ============================================================

# Account 799999 does not exist
# in the Trial Balance.

invoice_rows.append([
    "INV9998",
    "Unknown Supplier GmbH",
    "799999",
    "Austria",
    "Unmapped supplier invoice",
    12_500.00,
    "2025-12-31",
    "Posted"
])


# ============================================================
# 14. SAVE INVOICES
# ============================================================

invoice_file_path = (
    DATA_DIR / "Invoices_December_2025.xlsx"
)

invoice_workbook = Workbook()

invoice_worksheet = invoice_workbook.active
invoice_worksheet.title = "Invoices"


invoice_headers = [
    "Invoice ID",
    "Vendor",
    "Account",
    "Entity",
    "Description",
    "Amount",
    "Posting Date",
    "Status"
]

invoice_worksheet.append(
    invoice_headers
)

for row in invoice_rows:
    invoice_worksheet.append(row)


# Formatting

for cell in invoice_worksheet[1]:

    new_font = copy(cell.font)
    new_font.bold = True
    cell.font = new_font

invoice_worksheet.freeze_panes = "A2"
invoice_worksheet.auto_filter.ref = (
    invoice_worksheet.dimensions
)


invoice_column_widths = {
    "A": 14,
    "B": 30,
    "C": 14,
    "D": 15,
    "E": 30,
    "F": 16,
    "G": 16,
    "H": 12
}

for column, width in invoice_column_widths.items():
    invoice_worksheet.column_dimensions[column].width = width


invoice_workbook.save(
    invoice_file_path
)

print("Invoices successfully created.")
print(f"File: {invoice_file_path}")
print(f"Rows created: {len(invoice_rows)}")

# ============================================================
# 15. DATA VALIDATION CHECKS
# ============================================================

print("")
print("============================================================")
print("MONTH-END DATA VALIDATION")
print("============================================================")


# ------------------------------------------------------------
# CHECK 1 — TRIAL BALANCE BALANCE
# ------------------------------------------------------------

total_debit = sum(
    row[3]
    for row in tb_rows
)

total_credit = sum(
    row[4]
    for row in tb_rows
)

tb_difference = round(
    total_debit - total_credit,
    2
)

if tb_difference == 0:
    print("CHECK 1 - Trial Balance: PASS")
else:
    print(
        f"CHECK 1 - Trial Balance: FAIL "
        f"(Difference: {tb_difference:,.2f})"
    )


# ------------------------------------------------------------
# CHECK 2 — ACCRUAL ACCOUNT MAPPING
# ------------------------------------------------------------

valid_accounts = {
    account[0]
    for account in ACCOUNTS
}

unmapped_accruals = [
    row
    for row in accrual_rows
    if row[1] not in valid_accounts
]

if len(unmapped_accruals) == 0:
    print("CHECK 2 - Accrual Account Mapping: PASS")
else:
    print(
        "CHECK 2 - Accrual Account Mapping: "
        f"WARNING ({len(unmapped_accruals)} unmapped account)"
    )


# ------------------------------------------------------------
# CHECK 3 — PREPAYMENT ACCOUNT MAPPING
# ------------------------------------------------------------

unmapped_prepayments = [
    row
    for row in prepayment_rows
    if row[1] not in valid_accounts
    or row[2] not in valid_accounts
]

if len(unmapped_prepayments) == 0:
    print("CHECK 3 - Prepayment Account Mapping: PASS")
else:
    print(
        "CHECK 3 - Prepayment Account Mapping: "
        f"WARNING ({len(unmapped_prepayments)} unmapped account)"
    )


# ------------------------------------------------------------
# CHECK 4 — INVOICE ACCOUNT MAPPING
# ------------------------------------------------------------

unmapped_invoices = [
    row
    for row in invoice_rows
    if row[2] not in valid_accounts
]

if len(unmapped_invoices) == 0:
    print("CHECK 4 - Invoice Account Mapping: PASS")
else:
    print(
        "CHECK 4 - Invoice Account Mapping: "
        f"WARNING ({len(unmapped_invoices)} unmapped account)"
    )


# ------------------------------------------------------------
# CHECK 5 — DUPLICATE INVOICES
# ------------------------------------------------------------

invoice_keys = {}

for row in invoice_rows:

    key = (
        row[1],  # Vendor
        row[2],  # Account
        row[3],  # Entity
        row[5]   # Amount
    )

    invoice_keys[key] = (
        invoice_keys.get(key, 0) + 1
    )


duplicate_invoices = [
    key
    for key, count in invoice_keys.items()
    if count > 1
]

if len(duplicate_invoices) == 0:
    print("CHECK 5 - Duplicate Invoices: PASS")
else:
    print(
        "CHECK 5 - Duplicate Invoices: "
        f"WARNING ({len(duplicate_invoices)} duplicate)"
    )


# ------------------------------------------------------------
# CHECK 6 — SOURCE RECORD COUNTS
# ------------------------------------------------------------

print("")
print("SOURCE RECORD COUNTS")
print("---------------------")

print(f"Trial Balance : {len(tb_rows)}")
print(f"Accruals      : {len(accrual_rows)}")
print(f"Prepayments   : {len(prepayment_rows)}")
print(f"Invoices      : {len(invoice_rows)}")


# ------------------------------------------------------------
# FINAL STATUS
# ------------------------------------------------------------

print("")
print("============================================================")
print("DATA GENERATION AND VALIDATION COMPLETED")
print("============================================================")